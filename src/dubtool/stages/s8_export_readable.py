"""Stage 8: the customer-readable translation file (AGENT_UPDATE_7.md task 4).

One extra file next to every exported video, written for customers who are
not technical: the English original beside the Vietnamese, one row per
sentence, with the time the sentence is spoken IN THE EXPORTED VIDEO
(the `play_start` from schedule.json, not the original transcript start).

It reads only the small job files `translated.json`, `units.json` and
`schedule.json` and contains NO technical content: no ids, no model or
mode names, no token/timing numbers, no `fit_status`/`max_chars`/
`tts_duration`, none of the internal flags.

Rendering is failure-isolated by the caller (api.run_dub / cli.cmd_run):
any exception here never marks the video run failed (section 5.5).
"""

from __future__ import annotations

import logging
import os
from datetime import datetime
from pathlib import Path

log = logging.getLogger("dubtool")

DEDUPE_MIN_WORDS = 3  # same overlap rule as the s2/ASR cleanup (AGU3 7.3)


# ----------------------------------------------------------------- rows
def format_time(seconds: float) -> str:
    """mm:ss, or h:mm:ss from one hour (e.g. 1:02:03)."""
    total = max(0, int(round(float(seconds))))
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def _join_deduped(texts: list[str]) -> str:
    """Join a unit's segment texts, dropping words the recognizer repeated
    at a segment boundary (same overlap rule as dedupe_repeats in s2)."""
    from dubtool.stages.s2_transcribe import dedupe_repeats

    segs = [
        {"id": i, "text": t, "start": float(i), "end": float(i) + 1.0}
        for i, t in enumerate(texts)
    ]
    cleaned, _ = dedupe_repeats(segs, min_words=DEDUPE_MIN_WORDS)
    return " ".join(s["text"].strip() for s in cleaned if s["text"].strip())


def build_rows(job_dir: Path, mark_uncertain: bool = False) -> list[dict]:
    """One row per sentence: {"time": seconds, "en": text, "vi": text}.

    Sentence grouping comes from units.json; for older jobs without it the
    units are rebuilt from translated.json with the existing build_units.
    Rows with an empty English or Vietnamese text are skipped.

    `mark_uncertain` (AGU8 6.2, config readable.mark_uncertain): rows whose
    translation was still flagged after the completeness re-translation
    pass get the customer-facing "(cần kiểm tra)" marker on the Vietnamese
    text.
    """
    from dubtool.job import read_json
    from dubtool.schedule import load_schedule
    from dubtool.units import build_units

    entries = read_json(job_dir / "translated.json")
    if not isinstance(entries, list):
        raise RuntimeError("translated.json has an unexpected format")
    entries = sorted(entries, key=lambda e: int(e["id"]))
    by_id = {int(e["id"]): e for e in entries}

    units_path = job_dir / "units.json"
    if units_path.exists():
        units = read_json(units_path)
    else:
        units = build_units(
            [
                {"id": e["id"], "start": e["start"], "end": e["end"], "text": e.get("text") or ""}
                for e in entries
            ]
        )
    schedule = load_schedule(job_dir) or []
    play_start = {int(s["id"]): float(s["play_start"]) for s in schedule}

    rows: list[dict] = []
    for unit in units:
        seg_ids = [int(i) for i in unit.get("seg_ids", [])]
        texts_en = [by_id[i].get("text") or "" for i in seg_ids if i in by_id]
        texts_vi = [by_id[i].get("vi") or "" for i in seg_ids if i in by_id]
        en = _join_deduped([t for t in texts_en if t])
        vi = _join_deduped([t for t in texts_vi if t])
        if not en.strip() or not vi.strip():
            continue  # section 5.2: skip rows whose side is empty
        first = next((i for i in seg_ids if i in by_id), None)
        # Time in the EXPORTED video: the clip's play_start when the segment
        # took part in the schedule, else the transcript start of the unit.
        if first is not None and first in play_start:
            seconds = play_start[first]
        else:
            seconds = float(by_id[first]["start"]) if first is not None else 0.0
        if mark_uncertain and any(
            by_id[i].get("uncertain") for i in seg_ids if i in by_id
        ):
            vi = f"{vi} {place_uncertain()}".strip()
        rows.append({"time": seconds, "en": en, "vi": vi})
    return rows


# --------------------------------------------------------- unsure marker
def place_uncertain() -> str:
    """AGU8 6.2: the localized check marker for flagged translations."""
    from dubtool.app import i18n

    return str(i18n.tr("readable.uncertain") or "(cần kiểm tra)")


# ----------------------------------------------------------------- meta
def make_meta(video_out: Path, lang: str = "vi", duration_s: float | None = None) -> dict:
    """Header values for the document (no secrets, no technical fields)."""
    from dubtool.app import i18n

    return {
        "title": i18n.tr("readable.title"),
        "video": video_out.name,
        "date": datetime.now().strftime("%d/%m/%Y"),
        "duration": format_time(duration_s) if duration_s else "",
        "lang_pair": i18n.tr("readable.lang_pair"),
        "labels": {
            "video": i18n.tr("readable.info_video"),
            "duration": i18n.tr("readable.info_duration"),
            "lang": i18n.tr("readable.info_lang"),
            "date": i18n.tr("readable.info_created"),
            "col_time": i18n.tr("readable.col_time"),
            "col_en": i18n.tr("readable.col_en"),
            "col_vi": i18n.tr("readable.col_vi"),
        },
        "disclaimer": i18n.tr("readable.disclaimer"),
        "page_word": i18n.tr("readable.page", n="{PAGE}", t="{NUMPAGES}"),
    }


def readable_target(video_out: Path, fmt: str = "docx") -> Path:
    """`<video name>_vi_translation.docx|txt` next to the video.

    The number suffix travels with the video's own name, so the pair always
    keeps matching numbers (`clip_vi.mp4` -> `clip_vi_translation.docx`).
    """
    ext = ".docx" if str(fmt).lower() == "docx" else ".txt"
    return video_out.with_name(video_out.stem + "_translation" + ext)


def _atomic_write(path: Path, writer) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".part")
    writer(tmp)
    os.replace(tmp, path)


# ------------------------------------------------------------------ docx
_GREY = "595959"
_SOFT = "808080"
_HEADER_FILL = "DCE6F1"   # light blue-grey header band
_ALT_FILL = "F4F6F8"      # very light shading on every second row
_BORDER = "C9CFD6"        # thin light-grey borders


def render_docx(path: Path, meta: dict, rows: list[dict]) -> None:
    """The Word layout (section 5.3): A4 portrait, 2 cm margins, 20 pt bold
    title, grey-label info lines, a 3-column table with a repeating header
    row, rows that never split across pages, Arial 10.5 pt everywhere, and
    a footer with the disclaimer and Trang X / Y fields."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    def font_run(run, *, size=10.5, bold=False, color=None) -> None:
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = "Arial"
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        # Every script type, so Vietnamese diacritics render everywhere.
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(qn(attr), "Arial")

    def shade(cell, fill: str) -> None:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Cm(2.0))
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("000000")

    # Title.
    title = doc.add_paragraph()
    font_run(title.add_run(meta["title"]), size=20, bold=True)

    # Info lines: label in grey, value in bold.
    labels = meta["labels"]
    for key, value in (
        ("video", meta["video"]),
        ("duration", meta["duration"]),
        ("lang", meta["lang_pair"]),
        ("date", meta["date"]),
    ):
        if not value:
            continue  # e.g. a missing duration probe
        line = doc.add_paragraph()
        font_run(line.add_run(f"{labels[key]}: "), color=_GREY)
        font_run(line.add_run(value), bold=True)

    # The table: header row + one row per sentence.
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), _BORDER)
        borders.append(el)
    tblPr.append(borders)
    cell_mar = OxmlElement("w:tblCellMar")
    for edge, dxa in (("top", 80), ("bottom", 80), ("left", 110), ("right", 110)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(dxa))
        el.set(qn("w:type"), "dxa")
        cell_mar.append(el)
    tblPr.append(cell_mar)

    def row_flags(row, header: bool) -> None:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))  # a row never splits across pages
        if header:
            trPr.append(OxmlElement("w:tblHeader"))  # repeated on every page

    header = table.rows[0]
    row_flags(header, header=True)
    for cell, text in zip(header.cells, (labels["col_time"], labels["col_en"], labels["col_vi"])):
        shade(cell, _HEADER_FILL)
        para = cell.paragraphs[0]
        font_run(para.add_run(text), bold=True)
    for cell, width in zip(header.cells, (2.3, 7.35, 7.35)):
        cell.width = Cm(width)

    for idx, item in enumerate(rows):
        row = table.add_row()
        row_flags(row, header=False)
        cells = row.cells
        for cell, width in zip(cells, (2.3, 7.35, 7.35)):
            cell.width = Cm(width)
        t_para = cells[0].paragraphs[0]
        font_run(t_para.add_run(format_time(item["time"])), color=_GREY)
        for cell, text in ((cells[1], item["en"]), (cells[2], item["vi"])):
            para = cell.paragraphs[0]
            para.paragraph_format.line_spacing = 1.25
            font_run(para.add_run(text), size=10.5)
        if idx % 2 == 1:  # a very light shading on every second row
            for cell in cells:
                shade(cell, _ALT_FILL)

    # Footer on every page: disclaimer + Trang X / Y (Word fields).
    footer = section.footer
    p1 = footer.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p1.add_run(meta["disclaimer"]), size=8.5, color=_SOFT)
    p2 = footer.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_field(paragraph, instr: str) -> None:
        b = paragraph.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), "begin")
        b._element.append(el)
        f = paragraph.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        f._element.append(it)
        e = paragraph.add_run()
        el2 = OxmlElement("w:fldChar")
        el2.set(qn("w:fldCharType"), "end")
        e._element.append(el2)

    # "Trang {PAGE} / {NUMPAGES}" -> text, PAGE field, " / ", NUMPAGES.
    before, rest = meta["page_word"].split("{PAGE}", 1)
    between, after = (rest.split("{NUMPAGES}", 1) + [""])[:2]
    for part in (before, between, after):
        if part:
            font_run(p2.add_run(part), size=8.5, color=_SOFT)
        if part is before:
            add_field(p2, "PAGE")
        elif part is between:
            add_field(p2, "NUMPAGES")

    def save(out: Path) -> None:
        doc.save(str(out))

    _atomic_write(path, save)


# ------------------------------------------------------------------- txt
def render_txt(path: Path, meta: dict, rows: list[dict]) -> None:
    """The text fallback (section 5.4): UTF-8 BOM (Notepad shows the
    diacritics), the same title/info lines, then `[time] / EN: / VI:`
    blocks and the disclaimer at the end."""
    labels = meta["labels"]
    lines: list[str] = [meta["title"], ""]
    for key, value in (
        ("video", meta["video"]),
        ("duration", meta["duration"]),
        ("lang", meta["lang_pair"]),
        ("date", meta["date"]),
    ):
        if value:
            lines.append(f"{labels[key]}: {value}")
    lines.append("")
    for item in rows:
        lines.append(f"[{format_time(item['time'])}]")
        lines.append(f"EN: {item['en']}")
        lines.append(f"VI: {item['vi']}")
        lines.append("")
    lines.append(meta["disclaimer"])
    lines.append("")
    text = "\n".join(lines)

    def save(out: Path) -> None:
        # utf-8-sig writes the BOM so Windows Notepad shows diacritics.
        out.write_text(text, encoding="utf-8-sig", newline="\r\n")

    _atomic_write(path, save)


# --------------------------------------------------------------- delivery
def deliver(
    job_dir: Path,
    video_target: Path,
    *,
    fmt: str = "docx",
    lang: str = "vi",
    duration_s: float | None = None,
    mark_uncertain: bool = False,
) -> Path:
    """Build the readable file for the video delivered at `video_target`
    and return its path. Raises on failure (the CALLER isolates it; section
    5.5: the video result stays successful and the caller falls back to
    the .txt format)."""
    from dubtool.app import i18n

    i18n.set_language(lang)
    target = readable_target(video_target, fmt)
    meta = make_meta(video_target, lang=lang, duration_s=duration_s)
    rows = build_rows(job_dir, mark_uncertain=mark_uncertain)
    if not rows:
        raise RuntimeError("no sentence rows to export (empty translations?)")
    if str(fmt).lower() == "docx":
        render_docx(target, meta, rows)
    else:
        render_txt(target, meta, rows)
    log.info("s8: readable translation written to %s (%d rows)", target, len(rows))
    return target
